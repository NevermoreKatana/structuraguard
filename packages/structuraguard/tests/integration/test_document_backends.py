"""Реальные optional extras: стандартные packages и PDF text/table layer."""

from __future__ import annotations

import io
from typing import Protocol, cast

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from tests.contract_suites.parser import ParserContractCase, assert_parser_contract
from tests.unit.parsers.builtin._document_fixtures import pdf_bytes
from tests.unit.parsers.builtin._support import collect, contexts_for, source_for

from structuraguard.exceptions import ParserError, SecurityPolicyError
from structuraguard.parsers.builtin import DocxParser, PdfParser, XlsxParser

pytestmark = [pytest.mark.anyio, pytest.mark.integration]


class _WritablePdf(Protocol):
    """Узкий typed bridge к fixture-writing API native библиотеки."""

    def tobytes(self, *, encryption: int, owner_pw: str, user_pw: str) -> bytes: ...
    def embfile_add(self, name: str, buffer: bytes) -> object: ...
    def close(self) -> None: ...


class _WritablePdfBackend(Protocol):
    PDF_ENCRYPT_AES_256: int
    PDF_ENCRYPT_NONE: int

    def open(self, *, stream: bytes, filetype: str) -> _WritablePdf: ...


async def test_openpyxl_created_workbook() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Raw A", "Raw B"])
    sheet.append(["001", "=SUM(1,2)"])
    sheet.add_table(Table(displayName="PhysicalTable", ref="A1:B2"))
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    content = buffer.getvalue()
    source = source_for(content, display_name="real.xlsx")
    probe, parse = contexts_for(source, content, batch_size=1)
    batches = await assert_parser_contract(
        ParserContractCase(
            parser=XlsxParser(), source=source, probe_context=probe, parse_context=parse
        )
    )
    assert [
        cell.value.raw_value.value
        for batch in batches
        for table in batch.tables
        for cell in table.cells
    ] == ["Raw A", "Raw B", "001", ""]
    assert any(
        v.raw_value.value == "SUM(1,2)"
        for batch in batches
        for block in batch.blocks
        for v in block.values
    )


async def test_python_docx_created_document() -> None:
    document = Document()
    document.add_heading("Heading", level=1)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "001"
    document.add_paragraph("after table")
    buffer = io.BytesIO()
    document.save(buffer)
    content = buffer.getvalue()
    source = source_for(content, display_name="real.docx")
    probe, parse = contexts_for(source, content, batch_size=1)
    batches = await assert_parser_contract(
        ParserContractCase(
            parser=DocxParser(), source=source, probe_context=probe, parse_context=parse
        )
    )
    assert [
        b.text for batch in batches for b in batch.blocks if b.kind.value != "metadata"
    ] == ["Heading", "001", "after table"]


async def test_pdf_geometric_table_fixture() -> None:
    content = pdf_bytes(grid=True)
    source = source_for(content, display_name="grid.pdf")
    batches = await collect(PdfParser(), source, contexts_for(source, content)[1])
    table = batches[0].tables[0]
    assert [
        (c.row_index, c.column_index, c.value.raw_value.value) for c in table.cells
    ] == [(0, 0, "A"), (0, 1, "B"), (1, 0, "01"), (1, 1, "02")]
    assert table.location.kind == "document_block"


@pytest.mark.parametrize("encrypted", [True, False], ids=["encrypted", "embedded_file"])
async def test_pdf_encryption_and_real_embedded_file(encrypted: bool) -> None:
    backend = cast(_WritablePdfBackend, pymupdf)
    document = backend.open(stream=pdf_bytes(), filetype="pdf")
    try:
        if not encrypted:
            document.embfile_add("inert.txt", b"secret-canary")
        content = document.tobytes(
            encryption=backend.PDF_ENCRYPT_AES_256
            if encrypted
            else backend.PDF_ENCRYPT_NONE,
            owner_pw="owner",
            user_pw="user" if encrypted else "",
        )
    finally:
        document.close()
    source = source_for(content, display_name="protected.pdf")
    with pytest.raises(
        ParserError if encrypted else SecurityPolicyError,
        match="PARSER_UNSUPPORTED_FEATURE" if encrypted else "SECURITY_INPUT_REJECTED",
    ):
        await collect(PdfParser(), source, contexts_for(source, content)[1])
